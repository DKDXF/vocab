"""
AI Skills 相关 API —— 完形填空、作文优化、LLM 配置
"""
import json
import os
import uuid
from datetime import date, datetime
from typing import Optional
from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from database import get_db, close_db
from models import LLMSettingsOut, LLMSettingsUpdate

router = APIRouter(prefix="/api", tags=["AI Skills"])


def _get_setting(conn, key: str, default: str) -> str:
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    return row["value"] if row else default


def _get_llm_config(conn) -> dict:
    """获取 LLM 配置"""
    return {
        "api_key": _get_setting(conn, "llm_api_key", ""),
        "api_base": _get_setting(conn, "llm_api_base", "https://api.openai.com/v1"),
        "model": _get_setting(conn, "llm_model", "gpt-3.5-turbo"),
    }


def _check_llm_config(conn) -> str:
    """检查 LLM 配置是否完整，返回空字符串或错误信息"""
    config = _get_llm_config(conn)
    if not config["api_key"]:
        return "请先在设置中配置 AI API Key"
    return ""


# ==================== LLM 配置 ====================

@router.get("/settings/llm", response_model=LLMSettingsOut, summary="获取 AI 配置")
def get_llm_settings():
    """获取 AI 配置（key 脱敏显示）"""
    conn = get_db()
    try:
        api_key = _get_setting(conn, "llm_api_key", "")
        api_base = _get_setting(conn, "llm_api_base", "https://api.openai.com/v1")
        model = _get_setting(conn, "llm_model", "gpt-3.5-turbo")

        if api_key:
            masked = api_key[:8] + "..." + api_key[-4:] if len(api_key) > 12 else "***"
        else:
            masked = ""

        return LLMSettingsOut(
            api_key_set=bool(api_key),
            api_key_masked=masked,
            api_base=api_base,
            model=model,
        )
    finally:
        close_db(conn)


@router.post("/settings/llm", summary="更新 AI 配置")
def update_llm_settings(req: LLMSettingsUpdate):
    """更新 AI 配置"""
    conn = get_db()
    try:
        if req.api_key is not None:
            conn.execute(
                "INSERT INTO settings (key, value) VALUES ('llm_api_key', ?) "
                "ON CONFLICT(key) DO UPDATE SET value = ?",
                (req.api_key, req.api_key),
            )
        if req.api_base is not None:
            conn.execute(
                "INSERT INTO settings (key, value) VALUES ('llm_api_base', ?) "
                "ON CONFLICT(key) DO UPDATE SET value = ?",
                (req.api_base, req.api_base),
            )
        if req.model is not None:
            conn.execute(
                "INSERT INTO settings (key, value) VALUES ('llm_model', ?) "
                "ON CONFLICT(key) DO UPDATE SET value = ?",
                (req.model, req.model),
            )
        conn.commit()
        return {"message": "AI 配置已保存"}
    finally:
        close_db(conn)


# ==================== AI 完形填空 ====================

@router.post("/skills/cloze/generate", summary="生成完形填空")
async def generate_cloze(word_ids: Optional[str] = None):
    """
    基于用户学习的单词，调用大模型生成完形填空文章
    word_ids: 逗号分隔的单词ID，为空则用今日所学单词
    """
    conn = get_db()
    try:
        err = _check_llm_config(conn)
        if err:
            raise HTTPException(status_code=400, detail=err)

        book_id = _get_setting(conn, "current_book_id", "cet4")
        config = _get_llm_config(conn)

        # 获取单词列表
        if word_ids:
            ids = [int(x.strip()) for x in word_ids.split(",") if x.strip()]
            placeholders = ",".join("?" * len(ids))
            rows = conn.execute(
                f"SELECT * FROM words WHERE id IN ({placeholders})", ids
            ).fetchall()
        else:
            # 获取今日学习的单词
            today = date.today().isoformat()
            rows = conn.execute(
                """SELECT w.* FROM words w
                   JOIN study_records sr ON sr.word_id = w.id
                   WHERE w.book_id = ? AND sr.created_at LIKE ?
                   LIMIT 10""",
                (book_id, f"{today}%"),
            ).fetchall()
            if not rows:
                # 没有今日记录则取当前词书前10个
                rows = conn.execute(
                    "SELECT * FROM words WHERE book_id = ? LIMIT 10", (book_id,)
                ).fetchall()

        if not rows:
            raise HTTPException(status_code=400, detail="没有可用的单词")

        word_list = [{"word": r["word"], "definition": r["definition_cn"]} for r in rows]

        # 调用 AI 生成完形填空
        from skills.ai_cloze.skill import AIClozeSkill
        skill = AIClozeSkill(config["api_key"], config["api_base"], config["model"])
        result = await skill.execute(word_list=word_list)
        return result
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成失败: {str(e)}")
    finally:
        close_db(conn)


@router.post("/skills/cloze/submit", summary="提交完形填空答案")
async def submit_cloze(request_body: dict):
    """提交完形填空答案，返回得分"""
    answers = request_body.get("answers", {})
    cloze_data = request_body.get("cloze_data", {})
    blanks = cloze_data.get("blanks", [])

    correct = 0
    total = len(blanks)
    results = []

    for i, blank in enumerate(blanks):
        blank_idx = str(i)
        user_answer = answers.get(blank_idx, "")
        is_correct = user_answer.strip().lower() == blank.get("answer", "").strip().lower()
        if is_correct:
            correct += 1
        results.append({
            "index": i,
            "correct": is_correct,
            "user_answer": user_answer,
            "correct_answer": blank.get("answer", ""),
        })

    return {
        "total": total,
        "correct": correct,
        "score": round((correct / total) * 100) if total > 0 else 0,
        "results": results,
    }


# ==================== AI 作文优化 ====================

@router.post("/skills/writing/optimize", summary="AI 作文优化")
async def optimize_writing(
    mode: str = Form("text"),
    text: Optional[str] = Form(None),
    image: Optional[UploadFile] = File(None),
    file: Optional[UploadFile] = File(None),
):
    """提交作文进行 AI 优化（支持 text/image/file 三种 mode）"""
    conn = get_db()
    try:
        err = _check_llm_config(conn)
        if err:
            raise HTTPException(status_code=400, detail=err)

        config = _get_llm_config(conn)
        essay_text = text or ""

        if mode == "image" and image:
            # 图片模式：读取图片 base64，用多模态能力识别
            content = await image.read()
            import base64
            b64 = base64.b64encode(content).decode("utf-8")
            from skills.ai_writing.skill import AIWritingSkill
            skill = AIWritingSkill(config["api_key"], config["api_base"], config["model"])
            essay_text = await skill.extract_from_image(b64, image.content_type or "image/png")

        elif mode == "file" and file:
            # 文件模式：读取 txt/docx 文件
            content = await file.read()
            filename = file.filename or ""
            if filename.endswith(".txt"):
                essay_text = content.decode("utf-8")
            elif filename.endswith(".docx"):
                import io
                from docx import Document as DocxDocument
                doc = DocxDocument(io.BytesIO(content))
                essay_text = "\n".join([p.text for p in doc.paragraphs])
            else:
                raise HTTPException(status_code=400, detail="仅支持 .txt 和 .docx 文件")

        if not essay_text.strip():
            raise HTTPException(status_code=400, detail="作文内容不能为空")

        # 调用 AI 优化
        from skills.ai_writing.skill import AIWritingSkill
        skill = AIWritingSkill(config["api_key"], config["api_base"], config["model"])
        result = await skill.execute(text=essay_text)

        # 保存到历史
        today = datetime.now().isoformat()
        conn.execute(
            """INSERT INTO writing_history (original, optimized, feedback, score, mode, created_at)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (essay_text, result.get("optimized", ""), result.get("feedback", ""),
             result.get("score", 0), mode, today),
        )
        conn.commit()

        return result
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"优化失败: {str(e)}")
    finally:
        close_db(conn)


@router.get("/skills/writing/history", summary="获取作文优化历史")
def get_writing_history(limit: int = 20):
    """获取最近的作文优化历史"""
    conn = get_db()
    try:
        rows = conn.execute(
            "SELECT * FROM writing_history ORDER BY created_at DESC LIMIT ?",
            (limit,),
        ).fetchall()
        return [{
            "id": r["id"],
            "original": r["original"][:200] + "..." if len(r["original"]) > 200 else r["original"],
            "optimized": r["optimized"][:200] + "..." if len(r["optimized"]) > 200 else r["optimized"],
            "feedback": r["feedback"],
            "score": r["score"],
            "mode": r["mode"],
            "created_at": r["created_at"],
        } for r in rows]
    finally:
        close_db(conn)
